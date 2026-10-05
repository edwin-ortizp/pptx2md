import json
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import patch

from pptx import Presentation
from reportlab.pdfgen import canvas
from reportlab.lib.pdfencrypt import StandardEncryption

from pptx2md import ConversionConfig, convert
from pptx2md.__main__ import main


class PDFConversionTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)

    def pdf(self, name='source.pdf', encrypted=False, blank=False):
        path = self.root / name
        c = canvas.Canvas(str(path), encrypt=StandardEncryption('secret') if encrypted else None)
        if not blank:
            c.setFont('Helvetica', 18)
            c.drawString(72, 750, 'Definición de Scrum')
            c.setFont('Helvetica', 11)
            c.drawString(72, 720, 'Un párrafo con acentos que continúa')
            c.drawString(72, 704, 'en otra línea y también')
            c.drawString(540, 30, '1')
        c.showPage()
        if not blank:
            c.setFont('Helvetica', 11)
            c.drawString(72, 750, 'en la siguiente página.')
            c.drawString(90, 715, '1. Primera acción que continúa')
            c.drawString(104, 699, 'en otra línea.')
            c.drawString(90, 683, '2. Segunda acción.')
            c.drawString(90, 650, '- Una viñeta.')
            c.drawString(72, 610, 'Otro párrafo independiente.')
            c.drawString(540, 30, '2')
        c.save()
        return path

    def config(self, source, **kwargs):
        return ConversionConfig(pptx_path=source, output_path=self.root / 'result.md', image_dir=None, **kwargs)

    def test_unicode_lists_and_page_continuation(self):
        config = self.config(self.pdf())
        convert(config)
        text = config.output_path.read_text(encoding='utf-8')
        self.assertIn('# Definición de Scrum', text)
        self.assertIn('Un párrafo con acentos que continúa en otra línea y también en la siguiente página.', text)
        self.assertIn('1. Primera acción que continúa en otra línea.\n2. Segunda acción.', text)
        self.assertIn('* Una viñeta.', text)
        self.assertIn('\n\nOtro párrafo independiente.', text)
        self.assertNotIn('Slide ', text)
        self.assertNotIn('\n\n1\n', text)

    def test_page_selection_and_custom_titles(self):
        source = self.pdf()
        titles = self.root / 'titles.txt'
        titles.write_text('Documento\n  Definición de Scrum\n', encoding='utf-8')
        config = self.config(source, page=1, title_path=titles)
        convert(config)
        text = config.output_path.read_text(encoding='utf-8')
        self.assertIn('## Definición de Scrum', text)
        self.assertNotIn('siguiente página', text)
        with self.assertRaisesRegex(ValueError, 'between 1 and 2'):
            convert(self.config(source, page=3))

    def test_toc_wrapped_heading_and_numeric_body(self):
        source = self.root / 'toc.pdf'
        c = canvas.Canvas(str(source))
        c.setFont('Helvetica', 11)
        for i, line in enumerate(['Primera sección ........... 1', 'Segunda sección ........... 2',
                                  'Tercera sección ........... 3']):
            c.drawString(72, 750-i*25, line)
        c.showPage()
        c.setFont('Helvetica', 18)
        c.drawString(72, 750, 'Primera sección')
        c.drawString(72, 725, 'con título largo')
        c.setFont('Helvetica', 11)
        c.drawString(72, 685, '42')
        c.drawString(72, 650, 'Texto completo.')
        c.save()
        config = self.config(source)
        convert(config)
        text = config.output_path.read_text(encoding='utf-8')
        self.assertIn('* Primera sección\n* Segunda sección\n* Tercera sección', text)
        self.assertNotIn('.....', text)
        self.assertIn('# Primera sección con título largo', text)
        self.assertIn('42', text)

    def test_year_and_emphasis_are_prose_and_heading_levels_do_not_skip(self):
        source = self.root / 'hierarchy.pdf'
        c = canvas.Canvas(str(source))
        c.setFont('Helvetica', 18)
        c.drawString(72, 750, 'Main section')
        c.setFont('Helvetica', 13)
        c.drawString(72, 715, 'Subsection')
        c.setFont('Helvetica', 11)
        c.drawString(72, 685, 'Presentado en')
        c.drawString(72, 669, '1995. Continuación del párrafo.')
        c.setFont('Helvetica-Bold', 11)
        c.drawString(72, 635, 'Valores importantes')
        c.setFont('Helvetica', 11)
        c.drawString(72, 610, 'Texto normal.')
        c.drawString(72, 585, 'Información de contacto:')
        c.drawString(72, 569, 'Nombre: Ana')
        c.drawString(72, 553, 'Correo: ana@example.com')
        c.save()
        config = self.config(source)
        convert(config)
        text = config.output_path.read_text(encoding='utf-8')
        self.assertIn('## Subsection', text)
        self.assertIn('Presentado en 1995. Continuación del párrafo.', text)
        self.assertNotIn('# Valores importantes', text)
        self.assertIn('Información de contacto:\n\nNombre: Ana\n\nCorreo: ana@example.com', text)

    def test_errors_do_not_create_output(self):
        corrupt = self.root / 'corrupt.pdf'
        corrupt.write_bytes(b'not a pdf')
        for source, message in [(corrupt, 'could not read PDF'),
                                (self.pdf('locked.pdf', encrypted=True), 'password protected'),
                                (self.pdf('blank.pdf', blank=True), 'no extractable text')]:
            with self.subTest(source=source):
                config = self.config(source)
                with self.assertRaisesRegex(ValueError, message):
                    convert(config)
                self.assertFalse(config.output_path.exists())

    def test_json_preserves_ordered_markers(self):
        config = self.config(self.pdf())
        config.output_path = self.root / 'result.json'
        convert(config)
        elements = json.loads(config.output_path.read_text(encoding='utf-8'))['slides'][0]['elements']
        self.assertEqual([e['marker'] for e in elements if e['type'] == 'ListItem'], ['1.', '2.', None])

    def pptx(self, name='source.pptx'):
        source = self.root / name
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = 'PowerPoint title'
        slide.placeholders[1].text = 'Original PowerPoint paragraph.'
        prs.save(source)
        return source

    def test_pptx_regression(self):
        config = self.config(self.pptx())
        convert(config)
        self.assertEqual(config.output_path.read_text(encoding='utf-8'),
                         'Slide 1\n\n# PowerPoint title\n\nOriginal PowerPoint paragraph\\.\n\n')

    def test_batch_same_stem_and_uppercase_extension(self):
        self.pdf('guide.PDF')
        self.pptx('guide.pptx')
        out = self.root / 'output'
        with patch.object(sys, 'argv', ['pptx2md', str(self.root), '--all', '-o', str(out)]):
            main()
        self.assertEqual({p.name for p in out.iterdir()}, {'guide.PDF.md', 'guide.pptx.md'})
        self.assertIn('Un párrafo', (out / 'guide.PDF.md').read_text(encoding='utf-8'))
        self.assertIn('PowerPoint title', (out / 'guide.pptx.md').read_text(encoding='utf-8'))

    def test_batch_secondary_collision_rejected_before_writing(self):
        self.pdf('guide.pdf')
        self.pptx('guide.pptx')
        self.pdf('guide.pptx.pdf')
        out = self.root / 'output'
        with patch.object(sys, 'argv', ['pptx2md', str(self.root), '--all', '-o', str(out)]):
            with self.assertRaisesRegex(ValueError, 'collide'):
                main()
        self.assertFalse(out.exists())


if __name__ == '__main__':
    unittest.main()
