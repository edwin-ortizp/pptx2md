from pathlib import Path
import re
import tempfile
import unittest
from unittest.mock import patch
from urllib.parse import unquote

from PIL import Image
from pptx import Presentation
from pptx.util import Inches
from reportlab.pdfgen import canvas
from reportlab.lib.pdfencrypt import StandardEncryption

from pptx2md.batch import BatchOptions, combine_section, markdown_name, run_batch


class BatchTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)
        self.out = self.root / 'output'

    def pdf(self, name='document.pdf', encrypted=False):
        source = self.root / name
        source.parent.mkdir(parents=True, exist_ok=True)
        c = canvas.Canvas(str(source), encrypt=StandardEncryption('secret') if encrypted else None)
        c.setFont('Helvetica', 18)
        c.drawString(72, 750, 'Título del PDF')
        c.setFont('Helvetica', 11)
        c.drawString(72, 715, 'Este documento conserva el contenido completo y sus acentos.')
        c.save()
        return source

    def pptx(self, name='presentation.pptx'):
        source = self.root / name
        source.parent.mkdir(parents=True, exist_ok=True)
        image = self.root / 'picture.png'
        Image.new('RGB', (40, 40), 'blue').save(image)
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = 'Título PowerPoint'
        slide.placeholders[1].text = 'Texto de PowerPoint.'
        slide.shapes.add_picture(str(image), Inches(1), Inches(3), width=Inches(1))
        slide.notes_slide.notes_text_frame.text = 'Nota exclusiva del presentador'
        prs.save(source)
        return source

    def assert_no_temporary(self):
        self.assertFalse(list(self.out.glob('.pptx2md-*')))

    def assert_images(self, output):
        text = output.read_text(encoding='utf-8')
        links = re.findall(r'!\[\]\((.*?)\)', text)
        self.assertTrue(links)
        for link in links:
            self.assertNotIn('%5C', link)
            self.assertTrue((output.parent / unquote(link)).is_file(), link)

    def test_individual_mixed_progress_and_continued_errors(self):
        bad = self.root / 'broken.pdf'
        bad.write_bytes(b'broken')
        snapshots = []
        result = run_batch([self.pdf(), bad, self.pptx()], BatchOptions(self.out), snapshots.append)
        self.assertEqual([i['status'] for i in result['items']], ['completed', 'error', 'completed'])
        self.assertEqual((result['processed'], result['succeeded'], result['failed']), (3, 2, 1))
        self.assertEqual([s['current'] for s in snapshots if s['current'] and s['processed'] < s['total']
                          and s['items'][s['processed']]['status'] == 'converting'],
                         ['document.pdf', 'broken.pdf', 'presentation.pptx'])
        self.assertEqual(snapshots[0]['processed'], 0)
        self.assertFalse(snapshots[-1]['running'])
        self.assertNotIn('![](', Path(result['items'][2]['output']).read_text(encoding='utf-8'))
        self.assertNotIn('Nota exclusiva', Path(result['items'][2]['output']).read_text(encoding='utf-8'))
        self.assert_no_temporary()

    def test_combined_preserves_order_skips_failures_and_shifts_headings(self):
        pdf = self.pdf()
        locked = self.pdf('locked.pdf', encrypted=True)
        pptx = self.pptx()
        result = run_batch([pptx, locked, pdf], BatchOptions(self.out, 'combined'))
        output = Path(result['combined_output'])
        text = output.read_text(encoding='utf-8')
        self.assertLess(text.index('# presentation.pptx'), text.index('# document.pdf'))
        self.assertIn('## Título PowerPoint', text)
        self.assertIn('## Título del PDF', text)
        self.assertNotIn('# locked.pdf', text)
        self.assertIn('contraseña', result['items'][1]['error'])
        self.assertEqual(result['items'][0]['output'], str(output))
        self.assertEqual({p.name for p in self.out.glob('*.md')}, {'combinado.md'})
        self.assert_no_temporary()

    def test_images_and_notes_in_both_modes(self):
        first, second = self.pptx('one/lesson.pptx'), self.pptx('two/lesson.pptx')
        for mode in ('individual', 'combined'):
            with self.subTest(mode=mode):
                result = run_batch([first, second], BatchOptions(self.out, mode, include_images=True, include_notes=True))
                self.assertEqual(result['succeeded'], 2)
                files = {Path(i['output']) for i in result['items']}
                for output in files:
                    self.assert_images(output)
                    self.assertIn('Nota exclusiva', output.read_text(encoding='utf-8'))
                self.assert_no_temporary()

    def test_existing_names_duplicate_stems_and_duplicate_input(self):
        self.out.mkdir()
        existing = self.out / 'lesson.pptx.md'
        existing.write_text('Keep me', encoding='utf-8')
        a, b, c = self.pptx('lesson.pptx'), self.pdf('lesson.pdf'), self.pptx('other/lesson.pptx')
        result = run_batch([a, b, c, a], BatchOptions(self.out))
        self.assertEqual(result['total'], 3)
        self.assertEqual([Path(i['output']).name for i in result['items']],
                         ['lesson.pptx_2.md', 'lesson.pdf.md', 'lesson.pptx_3.md'])
        self.assertEqual(existing.read_text(encoding='utf-8'), 'Keep me')
        first = run_batch([b], BatchOptions(self.out, 'combined'))
        second = run_batch([b], BatchOptions(self.out, 'combined'))
        self.assertEqual(Path(first['combined_output']).name, 'combinado.md')
        self.assertEqual(Path(second['combined_output']).name, 'combinado_2.md')

    def test_all_failed_no_empty_combined_and_temp_cleanup(self):
        result = run_batch([self.pdf('locked.pdf', True)], BatchOptions(self.out, 'combined', include_images=True))
        self.assertIsNone(result['combined_output'])
        self.assertEqual(result['failed'], 1)
        self.assertFalse(list(self.out.glob('*.md')))
        self.assert_no_temporary()

    def test_final_combined_write_failure_invalidates_results_and_cleans_assets(self):
        with patch('pptx2md.batch._publish', side_effect=PermissionError('denied')):
            result = run_batch([self.pptx()], BatchOptions(self.out, 'combined', include_images=True))
        self.assertEqual(result['failed'], 1)
        self.assertEqual(result['succeeded'], 0)
        self.assertIsNone(result['items'][0]['output'])
        self.assertFalse((self.out / 'img' / 'combinado').exists())
        self.assert_no_temporary()

    def test_output_folder_failure_is_reported_for_every_file(self):
        self.out.write_text('Not a directory', encoding='utf-8')
        result = run_batch([self.pdf(), self.pptx()], BatchOptions(self.out))
        self.assertEqual(result['failed'], 2)
        self.assertIsNotNone(result['error'])
        self.assertFalse(result['running'])

    def test_publish_race_does_not_overwrite_existing_file(self):
        from pptx2md.batch import _publish
        source = self.root / 'temp.md'
        source.write_text('New', encoding='utf-8')
        target = self.root / 'existing.md'
        target.write_text('Old', encoding='utf-8')
        with self.assertRaises(FileExistsError):
            _publish(source, target)
        self.assertEqual(target.read_text(encoding='utf-8'), 'Old')

    def test_failed_conversion_cleans_partial_images_and_continues(self):
        sources = [self.pptx(), self.pdf()]
        from pptx2md.batch import convert as actual_convert
        def fail_first(config):
            if config.pptx_path.suffix == '.pptx':
                (config.image_dir / 'partial.png').write_bytes(b'partial')
                raise PermissionError('denied')
            actual_convert(config)
        with patch('pptx2md.batch.convert', side_effect=fail_first):
            result = run_batch(sources, BatchOptions(self.out, include_images=True))
        self.assertEqual(result['succeeded'], 1)
        self.assertFalse((self.out / 'img' / 'presentation').exists())
        self.assert_no_temporary()

    def test_asset_directory_race_does_not_delete_foreign_files(self):
        source = self.pptx()
        asset = self.out / 'img' / 'presentation'
        original = Path.mkdir
        def competing_mkdir(path, *args, **kwargs):
            if path == asset and not path.exists():
                original(path, parents=True)
                (path / 'existing.png').write_bytes(b'keep')
                raise FileExistsError('created by another process')
            return original(path, *args, **kwargs)
        with patch.object(Path, 'mkdir', competing_mkdir):
            result = run_batch([source], BatchOptions(self.out, include_images=True))
        self.assertEqual(result['failed'], 1)
        self.assertEqual((asset / 'existing.png').read_bytes(), b'keep')

    def test_heading_cap_fences_and_names(self):
        text = combine_section('doc[1].pdf', '###### Deep\n```\n# code\n```\n# Heading\n')
        self.assertIn('###### Deep', text)
        self.assertNotIn('#######', text)
        self.assertIn('\n# code\n', text)
        self.assertIn('\n## Heading\n', text)
        self.assertIn(r'doc\[1\].pdf', text)
        self.assertEqual(markdown_name('report'), 'report.md')
        for name in ['../outside.md', 'CON.md', 'report.pdf', '', 'bad?.md']:
            with self.assertRaises(ValueError):
                markdown_name(name)
